"""نتایج: خلاصه و نمودار هر سؤال، قیف ریزش، فیلتر، فهرست و جزئیات پاسخ‌ها، خروجی اکسل و CSV."""
import csv
import json
import statistics
from collections import Counter
from datetime import timedelta

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Avg, Count
from django.db.models.functions import TruncDate
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from apps.billing.utils import plan_of
from apps.common.fa import jdate, parse_jalali, to_fa
from . import engine
from .access import get_survey
from .models import Response


def _schema(s):
    return s.published or s.draft


def _filtered(request, s):
    qs = Response.objects.filter(survey=s, is_test=False)
    status = request.GET.get("status", "complete")
    if status == "complete":
        qs = qs.filter(completed_at__isnull=False)
    elif status == "partial":
        qs = qs.filter(completed_at__isnull=True)
    d1, d2 = parse_jalali(request.GET.get("from", "")), parse_jalali(request.GET.get("to", ""))
    if d1:
        qs = qs.filter(started_at__date__gte=d1)
    if d2:
        qs = qs.filter(started_at__date__lte=d2)
    fq, fv = request.GET.get("fq", ""), request.GET.get("fv", "")
    if fq and fv:
        qmap = {q["id"]: q for q in _schema(s).get("questions", [])}
        if fq in qmap:
            ids = [r.pk for r in qs.only("pk", "answers") if engine.eval_cond({"src": fq, "op": "eq", "value": fv}, r.answers, {}, {}, qmap)]
            qs = qs.filter(pk__in=ids)
    return qs, status


def _stats(q, answers_list):
    """آمار یک سؤال روی فهرست پاسخ‌ها."""
    t, o = q["type"], q.get("opt") or {}
    vals = [a.get(q["id"]) for a in answers_list]
    answered = [v for v in vals if not engine.is_empty(v)]
    st = {"q": q, "label": engine.TYPES[q["type"]][0], "answered": len(answered), "skipped": len(vals) - len(answered), "kind": "none"}
    n = len(answered) or 1
    if t in ("choice", "dropdown", "yes_no"):
        if t == "yes_no":
            opts = [("yes", o.get("yes_label", "بله")), ("no", o.get("no_label", "خیر"))]
        else:
            opts = [(c["id"], c["label"]) for c in o.get("choices", [])]
            if o.get("other"):
                opts.append(("__other__", o.get("other_label", "سایر")))
        cnt = Counter()
        for v in answered:
            for x in (v if isinstance(v, list) else [v]):
                cnt[x] += 1
        rows = [{"label": lbl, "n": cnt.get(k, 0), "pct": round(cnt.get(k, 0) * 100 / n)} for k, lbl in opts]
        st.update(kind="bars", rows=rows, chart={"labels": [r["label"] for r in rows], "values": [r["n"] for r in rows]},
                  pie=t == "yes_no" or len(rows) <= 4)
        if o.get("other"):
            st["others"] = [a.get(q["id"] + "__other") for a in answers_list if a.get(q["id"] + "__other")][:30]
    elif t in ("rating", "scale"):
        nums = [int(v) for v in answered if isinstance(v, (int, float))]
        lo = o.get("min", 1) if t == "scale" else 1
        hi = o.get("max", 10 if t == "scale" else 5)
        cnt = Counter(nums)
        rows = [{"label": to_fa(i), "n": cnt.get(i, 0), "pct": round(cnt.get(i, 0) * 100 / n)} for i in range(lo, hi + 1)]
        st.update(kind="bars", rows=rows, chart={"labels": [r["label"] for r in rows], "values": [r["n"] for r in rows]},
                  avg=round(statistics.mean(nums), 2) if nums else None, max=hi)
        if o.get("nps") and nums:
            pro = sum(1 for x in nums if x >= 9)
            det = sum(1 for x in nums if x <= 6)
            st["nps"] = {"score": round((pro - det) * 100 / len(nums)), "pro": round(pro * 100 / len(nums)),
                         "pas": round((len(nums) - pro - det) * 100 / len(nums)), "det": round(det * 100 / len(nums))}
    elif t in ("number", "price"):
        nums = [v for v in answered if isinstance(v, (int, float))]
        if nums:
            st.update(kind="numbers", avg=round(statistics.mean(nums), 2), median=statistics.median(nums), min=min(nums), max=max(nums),
                      total=sum(nums))
    elif t == "matrix":
        rows, cols = o.get("rows", []), o.get("cols", [])
        table = []
        for r in rows:
            cnt = Counter()
            for v in answered:
                c = (v or {}).get(r["id"]) if isinstance(v, dict) else None
                for x in (c if isinstance(c, list) else [c] if c else []):
                    cnt[x] += 1
            tot = sum(cnt.values()) or 1
            table.append({"label": r["label"], "cells": [{"n": cnt.get(c["id"], 0), "pct": round(cnt.get(c["id"], 0) * 100 / tot)} for c in cols]})
        st.update(kind="matrix", cols=[c["label"] for c in cols], table=table)
    elif t == "ranking":
        items = o.get("items", [])
        pos = {i["id"]: [] for i in items}
        for v in answered:
            for idx, x in enumerate(v or []):
                if x in pos:
                    pos[x].append(idx + 1)
        rows = sorted([{"label": i["label"], "avg": round(statistics.mean(pos[i["id"]]), 2) if pos[i["id"]] else None} for i in items],
                      key=lambda r: r["avg"] or 99)
        st.update(kind="ranking", rows=rows)
    elif t == "location":
        cnt = Counter((v or {}).get("province") for v in answered if isinstance(v, dict))
        rows = [{"label": k, "n": c, "pct": round(c * 100 / n)} for k, c in cnt.most_common(15)]
        st.update(kind="bars", rows=rows, chart={"labels": [r["label"] for r in rows], "values": [r["n"] for r in rows]})
    elif t in ("file", "signature"):
        st.update(kind="files", files=[v for v in answered if isinstance(v, dict)][-20:])
    elif t == "consent":
        st.update(kind="text", texts=[])
    else:
        st.update(kind="text", texts=[engine.display(q, v) for v in answered][-30:][::-1])
    return st


@login_required
def results(request, pk):
    s, m = get_survey(request, pk, "view")
    schema = _schema(s)
    qs, status = _filtered(request, s)
    rows = list(qs.values("answers", "variables", "completed_at", "duration", "device", "last_question"))
    answers_list = [r["answers"] for r in rows]
    stats = [_stats(q, answers_list) for q in schema.get("questions", []) if q["type"] in engine.ANSWERABLE]
    base = Response.objects.filter(survey=s, is_test=False)
    started, completed = base.count(), base.filter(completed_at__isnull=False).count()
    # قیف ریزش: آخرین سؤال پاسخ‌های نیمه‌کاره
    drop = Counter(base.filter(completed_at__isnull=True).values_list("last_question", flat=True))
    titles = {q["id"]: q["title"] for q in schema.get("questions", [])}
    funnel = [{"title": titles.get(k, "قبل از اولین سؤال"), "n": c} for k, c in drop.most_common(6)]
    # روند ۳۰ روزه
    today = timezone.localdate()
    days = [today - timedelta(days=i) for i in range(29, -1, -1)]
    daily = dict(base.filter(completed_at__isnull=False, completed_at__date__gte=days[0]).annotate(d=TruncDate("completed_at"))
                 .values_list("d").annotate(n=Count("id")))
    timeline = {"labels": [jdate(d, "%m/%d") for d in days], "values": [daily.get(d, 0) for d in days]}
    # متغیرها (مثل سطح درآمد یا امتیاز)
    var_stats = []
    for v in [{"name": "score", "label": "امتیاز"}] + schema.get("variables", []):
        vals = [r["variables"].get(v["name"]) for r in rows if r["variables"]]
        vals = [x for x in vals if x not in (None, "")]
        if not vals:
            continue
        if all(isinstance(x, (int, float)) for x in vals):
            if v["name"] == "score" and not any(vals):
                continue
            var_stats.append({"label": v.get("label", v["name"]), "kind": "num", "avg": round(statistics.mean(vals), 2),
                              "min": min(vals), "max": max(vals)})
        else:
            cnt = Counter(str(x) for x in vals)
            var_stats.append({"label": v.get("label", v["name"]), "kind": "cat",
                              "rows": [{"label": k, "n": c, "pct": round(c * 100 / len(vals))} for k, c in cnt.most_common(12)]})
    devices = Counter(r["device"] for r in rows)
    durs = [r["duration"] for r in rows if r["duration"]]
    filterable = [q for q in schema.get("questions", []) if q["type"] in ("choice", "dropdown", "yes_no")]
    return render(request, "surveys/results.html", {
        "s": s, "m": m, "tab": "results", "stats": stats, "n": len(rows), "started": started, "completed": completed,
        "rate": round(completed * 100 / started) if started else 0, "funnel": funnel, "timeline": json.dumps(timeline, ensure_ascii=False),
        "charts": json.dumps({str(i): st["chart"] for i, st in enumerate(stats) if st.get("chart")}, ensure_ascii=False),
        "var_stats": var_stats, "devices": [(k, c) for k, c in devices.most_common()], "avg_dur": round(statistics.median(durs)) if durs else 0,
        "status": status, "filterable": filterable, "plan": plan_of(s.workspace),
        "f": {"from": request.GET.get("from", ""), "to": request.GET.get("to", ""), "fq": request.GET.get("fq", ""), "fv": request.GET.get("fv", "")},
    })


@login_required
def responses(request, pk):
    s, m = get_survey(request, pk, "view")
    schema = _schema(s)
    qs, status = _filtered(request, s)
    q = request.GET.get("q", "").strip()
    if q:
        from django.db.models import TextField
        from django.db.models.functions import Cast
        qs = qs.annotate(_t=Cast("answers", TextField())).filter(_t__icontains=q) if not q.isdigit() else qs.filter(pk=int(q))
    page = Paginator(qs.order_by("-started_at"), 30).get_page(request.GET.get("page"))
    cols = [x for x in schema.get("questions", []) if x["type"] in engine.ANSWERABLE][:4]
    for r in page.object_list:
        r.cells = [engine.display(c, r.answers.get(c["id"]), r.answers)[:60] for c in cols]
    return render(request, "surveys/responses.html", {"s": s, "m": m, "tab": "responses", "page": page, "cols": cols, "status": status,
                                                      "q": q, "plan": plan_of(s.workspace)})


@login_required
def response_detail(request, pk, rid):
    s, m = get_survey(request, pk, "view")
    r = get_object_or_404(Response, pk=rid, survey=s)
    if request.method == "POST" and request.POST.get("act") == "delete":
        if not m or not m.can_edit:
            return redirect("surveys:response", pk=s.pk, rid=r.pk)
        if r.completed_at and not r.is_test:
            from django.db.models import F
            type(s).objects.filter(pk=s.pk, response_count__gt=0).update(response_count=F("response_count") - 1)
        r.delete()
        messages.success(request, "پاسخ حذف شد.")
        return redirect("surveys:responses", pk=s.pk)
    schema = _schema(s)
    items = []
    for q in schema.get("questions", []):
        if q["type"] not in engine.ANSWERABLE:
            continue
        v = r.answers.get(q["id"])
        items.append({"q": q, "value": engine.display(q, v, r.answers), "raw": v, "visited": q["id"] in (r.path or r.answers.keys())})
    var_labels = {"score": "امتیاز", **{v["name"]: v["label"] for v in schema.get("variables", [])}}
    ending = next((e for e in schema.get("endings", []) if e["id"] == r.ending), None)
    prev = Response.objects.filter(survey=s, is_test=False, pk__lt=r.pk).order_by("-pk").values_list("pk", flat=True).first()
    nxt = Response.objects.filter(survey=s, is_test=False, pk__gt=r.pk).order_by("pk").values_list("pk", flat=True).first()
    return render(request, "surveys/response_detail.html", {
        "s": s, "m": m, "tab": "responses", "r": r, "items": items, "ending": ending, "prev": prev, "next": nxt,
        "vars": [(var_labels.get(k, k), v) for k, v in (r.variables or {}).items() if not (k == "score" and not v)],
    })


def _export_rows(s, qs):
    schema = _schema(s)
    qcols = [q for q in schema.get("questions", []) if q["type"] in engine.ANSWERABLE]
    var_names = ["score"] + [v["name"] for v in schema.get("variables", [])]
    var_labels = {"score": "امتیاز", **{v["name"]: v["label"] for v in schema.get("variables", [])}}
    ends = {e["id"]: e["title"] for e in schema.get("endings", [])}
    head = ["شماره", "وضعیت", "شروع", "تکمیل", "مدت (ثانیه)", "دستگاه", "موبایل تأییدشده"] + [q["title"][:120] or q["id"] for q in qcols] \
        + [var_labels[v] for v in var_names] + [f"[{h}]" for h in s.hidden_list] + ["صفحه‌ی پایان"]
    yield head
    for r in qs.order_by("pk").iterator():
        yield ([r.pk, "کامل" if r.completed_at else "نیمه‌کاره", jdate(r.started_at, "%Y/%m/%d %H:%M"),
                jdate(r.completed_at, "%Y/%m/%d %H:%M") if r.completed_at else "", r.duration, r.device, r.mobile]
               + [engine.display(q, r.answers.get(q["id"]), r.answers) for q in qcols]
               + [r.variables.get(v, "") for v in var_names] + [r.hidden.get(h, "") for h in s.hidden_list] + [ends.get(r.ending, "")])


@login_required
def export(request, pk, fmt):
    s, m = get_survey(request, pk, "view")
    qs, _ = _filtered(request, s)
    stamp = jdate(timezone.now(), "%Y-%m-%d").translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789"))
    if fmt == "xlsx":
        if not plan_of(s.workspace).export:
            messages.error(request, "خروجی اکسل در پلن فعلی نیست؛ از CSV استفاده کنید یا پلن را ارتقا دهید.")
            return redirect("surveys:responses", pk=s.pk)
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "پاسخ‌ها"
        ws.sheet_view.rightToLeft = True
        for i, row in enumerate(_export_rows(s, qs)):
            ws.append([str(x) if isinstance(x, (dict, list)) else x for x in row])
            if i == 0:
                for c in ws[1]:
                    c.font = Font(bold=True, color="FFFFFF")
                    c.fill = PatternFill("solid", fgColor="1E9E7B")
                    c.alignment = Alignment(wrap_text=True, vertical="center")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 22
        ws.freeze_panes = "B2"
        from io import BytesIO
        buf = BytesIO()
        wb.save(buf)
        resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="responses-{s.code}-{stamp}.xlsx"'
        return resp
    resp = HttpResponse(content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="responses-{s.code}-{stamp}.csv"'
    resp.write("﻿")
    w = csv.writer(resp)
    for row in _export_rows(s, qs):
        w.writerow(row)
    return resp


@login_required
@require_POST
def clear_tests(request, pk):
    s, m = get_survey(request, pk, "edit")
    n = Response.objects.filter(survey=s, is_test=True).delete()[0]
    messages.success(request, f"{to_fa(n)} پاسخ آزمایشی پاک شد.")
    return redirect("surveys:responses", pk=s.pk)
